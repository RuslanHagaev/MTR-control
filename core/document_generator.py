from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def value(value):
    """
    Возвращает безопасное значение для вставки в документ.
    """
    if value is None:
        return ""

    return str(value)


def generate_material_act(act):
    """
    Формирует редактируемый DOCX-файл АоРПИ
    на основании данных MaterialAct.
    """

    document = Document()

    # =========================================================
    # НАСТРОЙКИ ДОКУМЕНТА
    # =========================================================

    section = document.sections[0]

    section.top_margin = Pt(28)
    section.bottom_margin = Pt(28)
    section.left_margin = Pt(35)
    section.right_margin = Pt(25)

    # =========================================================
    # ШАПКА
    # =========================================================

    table = document.add_table(
        rows=2,
        cols=2
    )

    table.style = "Table Grid"

    table.cell(0, 0).text = (
        f"Заказчик: {value(act.customer)}"
    )

    table.cell(0, 1).text = (
        "Основание:\n"
        "Форма №3 (рекомендуемая)\n"
        "ВСН 012-88 (часть II)\n"
        "Миннефтегазстрой"
    )

    table.cell(1, 0).text = (
        f"Подрядчик: {value(act.contractor)}"
    )

    table.cell(1, 1).text = (
        f"Объект: {value(act.object_name)}"
    )

    # =========================================================
    # ЗАГОЛОВОК
    # =========================================================

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        f"\nАКТ № {value(act.number)}\n"
        "о результатах проверки изделий"
    )

    run.bold = True
    run.font.size = Pt(14)

    # =========================================================
    # ДАТА
    # =========================================================

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph.add_run(
        f"от {value(act.date)}"
    )

    # =========================================================
    # ОБЩИЕ ДАННЫЕ
    # =========================================================

    table = document.add_table(
        rows=0,
        cols=2
    )

    table.style = "Table Grid"

    rows = [
        ("Вид изделий", act.product_type),
        ("Вид проверки", act.get_inspection_type_display()),
        ("Наименование изделий", act.product_name),
        ("Объект", act.object_name),
        ("Позиция", act.position_name),
        ("Шифр проекта", act.project_code),
        ("Проект / чертёж", act.project_document),
        (
            "Дата проекта / чертежа",
            act.project_document_date
        ),
        ("Рабочие чертежи", act.working_drawings),
        (
            "Участок / привязка / км / ПК",
            act.construction_section
        ),
    ]

    for label, field_value in rows:

        row = table.add_row()

        row.cells[0].text = label
        row.cells[1].text = value(field_value)

    # =========================================================
    # ТЕКСТ АКТА
    # =========================================================

    document.add_paragraph()

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "Составлен представителями:"
    ).bold = True

    # =========================================================
    # ПРЕДСТАВИТЕЛЬ ПОДРЯДЧИКА
    # =========================================================

    document.add_paragraph(
        "Строительной организации"
    )

    table = document.add_table(
        rows=2,
        cols=2
    )

    table.style = "Table Grid"

    table.cell(0, 0).text = "Должность"
    table.cell(0, 1).text = value(
        act.contractor_position
    )

    table.cell(1, 0).text = "Фамилия, инициалы"
    table.cell(1, 1).text = value(
        act.contractor_representative
    )

    # =========================================================
    # ПРЕДСТАВИТЕЛЬ ЗАКАЗЧИКА
    # =========================================================

    document.add_paragraph(
        "Заказчика"
    )

    table = document.add_table(
        rows=2,
        cols=2
    )

    table.style = "Table Grid"

    table.cell(0, 0).text = "Должность"
    table.cell(0, 1).text = value(
        act.customer_position
    )

    table.cell(1, 0).text = "Фамилия, инициалы"
    table.cell(1, 1).text = value(
        act.customer_representative
    )

    # =========================================================
    # РЕЗУЛЬТАТЫ ПРОВЕРКИ
    # =========================================================

    document.add_paragraph()

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "1. Осмотр геометрических размеров и маркировки"
    ).bold = True

    table = document.add_table(
        rows=0,
        cols=2
    )

    table.style = "Table Grid"

    inspection_rows = [
        (
            "Геометрические размеры",
            act.geometric_dimensions
        ),
        (
            "Маркировка",
            act.marking
        ),
        (
            "Технические условия",
            act.technical_conditions
        ),
        (
            "Соответствие",
            act.compliance_result
        ),
        (
            "Рабочие чертежи",
            act.working_drawings
        ),
    ]

    for label, field_value in inspection_rows:

        row = table.add_row()

        row.cells[0].text = label
        row.cells[1].text = value(field_value)

    # =========================================================
    # СОПРОВОДИТЕЛЬНАЯ ДОКУМЕНТАЦИЯ
    # =========================================================

    document.add_paragraph()

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "2. Сопроводительная документация"
    ).bold = True

    document.add_paragraph(
        value(act.accompanying_documents)
    )

    document.add_paragraph(
        "Документация в полном комплекте: "
        + (
            "ДА"
            if act.documents_complete
            else "НЕТ"
        )
    )

    # =========================================================
    # МЕХАНИЧЕСКИЕ СВОЙСТВА
    # =========================================================

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "3. Характеристики механических свойств"
    ).bold = True

    document.add_paragraph(
        value(act.mechanical_properties)
    )

    document.add_paragraph(
        value(act.mechanical_properties_result)
    )

    # =========================================================
    # ТРЕБОВАНИЯ
    # =========================================================

    table = document.add_table(
        rows=0,
        cols=2
    )

    table.style = "Table Grid"

    requirements = [
        (
            "Требования проекта",
            act.project_requirements
        ),
        (
            "Требования технических условий",
            act.technical_requirements
        ),
        (
            "Заключение",
            act.conclusion
        ),
    ]

    for label, field_value in requirements:

        row = table.add_row()

        row.cells[0].text = label
        row.cells[1].text = value(field_value)

    # =========================================================
    # ПОДПИСИ
    # =========================================================

    document.add_paragraph()

    paragraph = document.add_paragraph()

    paragraph.add_run(
        "Представитель:"
    ).bold = True

    table = document.add_table(
        rows=3,
        cols=4
    )

    table.style = "Table Grid"

    table.cell(0, 0).text = "Сторона"
    table.cell(0, 1).text = "Фамилия, инициалы"
    table.cell(0, 2).text = "Подпись"
    table.cell(0, 3).text = "Дата"

    table.cell(1, 0).text = "Подрядчик"
    table.cell(1, 1).text = value(
        act.contractor_representative
    )
    table.cell(1, 2).text = ""
    table.cell(1, 3).text = value(
        act.contractor_signature_date
    )

    table.cell(2, 0).text = "Заказчик"
    table.cell(2, 1).text = value(
        act.customer_representative
    )
    table.cell(2, 2).text = ""
    table.cell(2, 3).text = value(
        act.customer_signature_date
    )

    # =========================================================
    # ШРИФТ
    # =========================================================

    for paragraph in document.paragraphs:

        for run in paragraph.runs:

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    for run in paragraph.runs:

                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)

    # =========================================================
    # СОХРАНЕНИЕ В ПАМЯТЬ
    # =========================================================

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output
